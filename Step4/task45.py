from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.filters import Filters, FilterColumn, AutoFilter



wb = Workbook()
ws = wb.active
ws.title = "Users"

ws.append(["Name", "Age","City"])
users = [
    ["Serhii", 24, "Torun"],
    ["Alex", 30, "Bydgoszcz"],
    ["John", 18, "Warszawa"],
    ["Mike", 27, "Gdansk"],
]
for user in users:
    ws.append(user)

wb.create_sheet("Users Sorted")
ws_users_sorted = wb["Users Sorted"]
ws_users_sorted.append(["Name", "Age", "City"])

users_sorted = users.copy()

users_sorted.sort(key=lambda x: x[1])

for user in users_sorted:
    ws_users_sorted.append(user)

table = Table(
    displayName="UsersTable",
    ref=ws.calculate_dimension(),

)
table.autoFilter = AutoFilter(
    ref=ws.calculate_dimension(),
)
table_users_sorted = Table(
    displayName="UsersTableSorted",
    ref=ws_users_sorted.calculate_dimension(),
)



filters = Filters(filter=["Torun"])

filter_column = FilterColumn(
    colId= 2,
    filters=filters
)

table.autoFilter.filterColumn = [filter_column]

ws.add_table(table)
ws_users_sorted.add_table(table_users_sorted)
wb.save("users_sort_report_filter_test.xlsx")